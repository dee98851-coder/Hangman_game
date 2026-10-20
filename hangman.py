
import random
from pathlib import Path

try:
    from docx import Document
except ImportError:
    Document = None


# ==================== SETTINGS ====================

WORD_FILE = Path("hangman_words.docx").with_name("hangman_words.docx")
MAX_ATTEMPTS = 6

HANGMAN_ART = [
    r"""
     +---+
     |   |
         |
         |
         |
    =========
    """,
    r"""
     +---+
     |   |
     O   |
         |
         |
    =========
    """,
    r"""
     +---+
     |   |
     O   |
     |   |
         |
    =========
    """,
    r"""
     +---+
     |   |
     O   |
    /|   |
         |
    =========
    """,
    r"""
     +---+
     |   |
     O   |
    /|\  |
         |
    =========
    """,
    r"""
     +---+
     |   |
     O   |
    /|\  |
    /    |
    =========
    """,
    r"""
     +---+
     |   |
     O   |
    /|\  |
    / \  |
    =========
    """
]


# ==================== LOAD WORDS ====================

def load_words():
    if Document is None:
        print("\nERROR: The python-docx library is missing.")
        print("Install it using: python -m pip install python-docx")
        return []

    if not WORD_FILE.is_file():
        print(f"\nERROR: File not found: {WORD_FILE.name}")
        print("Place hangman_words.docx in the same folder as hangman.py.")
        return []

    try:
        document = Document(str(WORD_FILE))
        words = set()

        # Read words from paragraphs.
        texts = [paragraph.text for paragraph in document.paragraphs]

        # Read words from tables.
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    texts.append(cell.text)

        # Extract unique alphabetic words.
        for text in texts:
            for item in text.split():
                word = "".join(
                    character for character in item.upper()
                    if character.isalpha()
                )

                if len(word) >= 2:
                    words.add(word)

        return sorted(words)

    except Exception as error:
        print(f"\nERROR: Could not read the Word document: {error}")
        return []


# ==================== DISPLAY GAME ====================

def display_game(secret_word, guessed_letters, attempts_left):
    mistakes = MAX_ATTEMPTS - attempts_left

    print("\n" + "=" * 42)
    print("              H A N G M A N")
    print("=" * 42)

    print(HANGMAN_ART[mistakes])

    masked_word = " ".join(
        letter if letter in guessed_letters else "_"
        for letter in secret_word
    )

    print("WORD:     ", masked_word)
    print(f"ATTEMPTS:  {attempts_left}/{MAX_ATTEMPTS}")

    if guessed_letters:
        print("GUESSED:  ", " ".join(sorted(guessed_letters)))
    else:
        print("GUESSED:   None")

    print("-" * 42)


# ==================== PLAY GAME ====================

def play_game(words):
    secret_word = random.choice(words)
    guessed_letters = set()
    attempts_left = MAX_ATTEMPTS

    print("\nYour secret word is ready!")
    print("Guess one letter at a time.")

    while attempts_left > 0:
        # Victory check.
        if all(letter in guessed_letters for letter in secret_word):
            display_game(secret_word, guessed_letters, attempts_left)

            print("\n" + "*" * 42)
            print("             YOU WIN!")
            print("*" * 42)
            print(f"Congratulations! You guessed: {secret_word}")
            print("You saved the stick figure!")
            print("*" * 42)
            return True

        display_game(secret_word, guessed_letters, attempts_left)

        guess = input("Enter a letter (or type QUIT to exit): ")
        guess = guess.strip().upper()

        if guess == "QUIT":
            print("\nYou left the current game.")
            print(f"The secret word was: {secret_word}")
            return None

        if len(guess) != 1 or not guess.isalpha():
            print("Invalid input. Enter exactly one letter.")
            continue

        if guess in guessed_letters:
            print(f"You already guessed '{guess}'. Try another letter.")
            continue

        guessed_letters.add(guess)

        if guess in secret_word:
            print(f"\nCORRECT! '{guess}' is in the word.")
        else:
            attempts_left -= 1
            print(f"\nWRONG! '{guess}' is not in the word.")
            print(f"You have {attempts_left} incorrect guesses remaining.")

    # Defeat screen.
    display_game(secret_word, guessed_letters, attempts_left)

    print("\n" + "*" * 42)
    print("             GAME OVER!")
    print("*" * 42)
    print(f"The correct word was: {secret_word}")
    print("Better luck next time!")
    print("*" * 42)

    return False


# ==================== MAIN PROGRAM ====================

def main():
    print("=" * 42)
    print("          WELCOME TO HANGMAN")
    print("=" * 42)

    words = load_words()

    if not words:
        print("\nGame could not start. Fix the error above and try again.")
        return

    print(f"\nSuccessfully loaded {len(words)} unique words.")
    print(f"Maximum incorrect guesses: {MAX_ATTEMPTS}")
    print("\nRules:")
    print("- Guess the hidden word one letter at a time.")
    print("- Each incorrect guess adds a body part.")
    print("- You lose after six incorrect guesses.")
    print("- Type QUIT during a game to leave it.")

    wins = 0
    losses = 0
    games_played = 0

    while True:
        result = play_game(words)

        if result is None:
            break

        games_played += 1

        if result:
            wins += 1
        else:
            losses += 1

        print("\n" + "-" * 42)
        print("              SCOREBOARD")
        print("-" * 42)
        print(f"Games played: {games_played}")
        print(f"Wins:         {wins}")
        print(f"Losses:       {losses}")
        print("-" * 42)

        again = input("\nWould you like to play again? (y/n): ")
        if again.strip().lower() not in ("y", "yes"):
            break

    print("\n" + "=" * 42)
    print("           THANKS FOR PLAYING!")
    print("=" * 42)
    print(f"Final score: {wins} wins, {losses} losses.")
    print("Goodbye!")


if __name__ == "__main__":
    main()
